"""Tests for Excel merged-cell reader."""

import io

import pytest
from openpyxl import Workbook

from app.services.excel_reader import detect_header_row, read_excel_with_merges


def _make_merged_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"

    ws["A1"] = "Region"
    ws["B1"] = "Revenue"
    ws.merge_cells("A2:A4")
    ws["A2"] = "North"
    ws["B2"] = 100
    ws["B3"] = 200
    ws["B4"] = 300

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_resolve_merged_data_cells():
    content = _make_merged_xlsx()
    df, meta = read_excel_with_merges(content)

    assert meta.merged_cells_resolved == 1
    assert meta.merge_support is True
    assert meta.sheet_name == "Sales"
    assert list(df["Region"]) == ["North", "North", "North"]
    assert list(df["Revenue"]) == [100, 200, 300]


def test_detect_header_row():
    grid = [
        ["Report Title", None],
        ["Region", "Revenue"],
        ["North", 100],
    ]
    assert detect_header_row(grid) == 1


def test_custom_header_row():
    content = _make_merged_xlsx()
    df, meta = read_excel_with_merges(content, header_row=0)

    assert meta.header_row == 0
    assert "Region" in df.columns or "Revenue" in df.columns
