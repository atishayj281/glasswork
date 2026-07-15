"""Read Excel files with merged-cell resolution via openpyxl."""

import io
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from app.config import MAX_ROWS
from app.models.schema import ExcelIngestMeta


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def _build_grid(ws) -> list[list[Any]]:
    max_row = min(ws.max_row or 0, MAX_ROWS + 50)
    max_col = ws.max_column or 0
    grid: list[list[Any]] = [
        [_cell_value(ws.cell(row=r, column=c).value) for c in range(1, max_col + 1)]
        for r in range(1, max_row + 1)
    ]
    return grid


def _resolve_merged_cells(ws, grid: list[list[Any]]) -> int:
    count = 0
    for merged_range in ws.merged_cells.ranges:
        min_row, min_col = merged_range.min_row, merged_range.min_col
        max_row, max_col = merged_range.max_row, merged_range.max_col
        value = ws.cell(row=min_row, column=min_col).value
        value = _cell_value(value)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                grid[r - 1][c - 1] = value
        count += 1
    return count


def detect_header_row(grid: list[list[Any]]) -> int:
    """Pick the row that looks most like a header (strings, followed by data)."""
    best_row = 0
    best_score = -1.0

    for i, row in enumerate(grid[:20]):
        non_empty = [v for v in row if v is not None and str(v).strip()]
        if not non_empty:
            continue
        str_count = sum(1 for v in non_empty if isinstance(v, str))
        score = str_count / len(non_empty)

        if i + 1 < len(grid):
            next_row = grid[i + 1]
            next_vals = [v for v in next_row if v is not None]
            if next_vals:
                num_count = sum(1 for v in next_vals if isinstance(v, (int, float)))
                if num_count > 0:
                    score += 0.5

        if score >= 0.5 and score > best_score:
            best_score = score
            best_row = i

    return best_row


def _clean_column_names(raw_names: list[Any]) -> list[str]:
    names: list[str] = []
    last_valid = ""

    for i, raw in enumerate(raw_names):
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            name = last_valid if last_valid else f"column_{i + 1}"
        else:
            name = str(raw).strip()
            last_valid = name

        if re.match(r"^Unnamed:\s*\d+$", name, re.I):
            name = last_valid if last_valid else f"column_{i + 1}"

        names.append(name)

    seen: dict[str, int] = {}
    deduped: list[str] = []
    for name in names:
        if name in seen:
            seen[name] += 1
            deduped.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            deduped.append(name)

    return deduped


def _drop_empty_edges(df: pd.DataFrame) -> pd.DataFrame:
    df = df.dropna(how="all")
    df = df.loc[:, df.notna().any()]
    return df.reset_index(drop=True)


def read_excel_with_merges(
    content: bytes,
    *,
    sheet_index: int = 0,
    header_row: int | None = None,
    max_rows: int = MAX_ROWS,
) -> tuple[pd.DataFrame, ExcelIngestMeta]:
    wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    sheets = wb.worksheets
    if sheet_index >= len(sheets):
        raise ValueError(f"Sheet index {sheet_index} out of range (file has {len(sheets)} sheets)")

    ws = sheets[sheet_index]
    grid = _build_grid(ws)
    merged_count = _resolve_merged_cells(ws, grid)

    if not grid:
        meta = ExcelIngestMeta(
            sheet_name=ws.title,
            header_row=0,
            merged_cells_resolved=merged_count,
            merge_support=True,
        )
        return pd.DataFrame(), meta

    hdr = header_row if header_row is not None else detect_header_row(grid)
    if hdr >= len(grid):
        hdr = 0

    raw_headers = grid[hdr]
    columns = _clean_column_names(raw_headers)
    data_rows = grid[hdr + 1 : hdr + 1 + max_rows]

    df = pd.DataFrame(data_rows, columns=columns)
    df = _drop_empty_edges(df)

    meta = ExcelIngestMeta(
        sheet_name=ws.title,
        header_row=hdr,
        merged_cells_resolved=merged_count,
        merge_support=True,
    )
    return df, meta
